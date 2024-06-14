import configparser
from PyQt5.QtWidgets import QTableWidgetItem, QGroupBox, QVBoxLayout, QPushButton, QTextEdit, QSizePolicy, \
     QSpacerItem, QTableWidget, QHeaderView, QFileDialog, QLineEdit, QDialog, QCheckBox, QScrollArea, QProgressBar
from PyQt5.QtCore import Qt, QSize, pyqtSignal
from gui.frontend.utils import create_transparent_button, setup_button_style
import pandas as pd
from fpdf import FPDF
import os
import webbrowser
from gui.frontend.settings_window_GUI import SettingsWindowGUI
from gui.frontend.help_window_GUi import helpWindowGui
from functools import partial
import json
from PyQt5.QtGui import QMovie, QCursor


try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

from PyQt5.QtWidgets import QLabel, QHBoxLayout, QWidget


class CustomTableWidgetItem(QTableWidgetItem):
    def __lt__(self, other):
        try:
            return int(self.text()) < int(other.text())
        except ValueError:
            return self.text() < other.text()

def display_output(self, headers, data):
    # display the output data in the table
    print("Displaying output in table")
    self.headers = headers
    self.data = data

    self.outputTable.setColumnCount(len(headers))
    self.outputTable.setHorizontalHeaderLabels(headers)
    self.outputTable.setRowCount(len(data))

    for row_idx, row_data in enumerate(data):
        for col_idx, col_data in enumerate(row_data):
            if isinstance(col_data, QMovie):
                # QLabel to display the loading wheel
                label = QLabel()
                label.setAlignment(Qt.AlignCenter)
                label.setMovie(col_data)
                col_data.start()  # Start the loading wheel

                # create a QWidget to hold the QLabel
                cell_widget = QWidget()
                layout = QHBoxLayout(cell_widget)
                layout.addWidget(label)
                layout.setAlignment(Qt.AlignCenter)
                layout.setContentsMargins(0, 0, 0, 0)
                cell_widget.setLayout(layout)

                self.outputTable.setCellWidget(row_idx, col_idx, cell_widget)
            else:
                item = CustomTableWidgetItem(str(col_data))
                item.setTextAlignment(Qt.AlignCenter)
                self.outputTable.setItem(row_idx, col_idx, item)

    for col in range(len(headers)):
        self.sort_orders[col] = Qt.DescendingOrder




class RightGroupBox(QGroupBox):
    back_to_home_signal = pyqtSignal()
    row_selected_signal = pyqtSignal(list)
    pid_selected_signal = pyqtSignal(str)

    def __init__(self, parent):
        super().__init__(parent)
        self.setObjectName("groupBox_right")
        self.setStyleSheet("QWidget { background-color: #262626; }")
        self.setFlat(True)
        self.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        self.initialize_ui()
        self.sort_orders = {}
        self.headers = []
        self.data = []
        self.filter_settings = {"all_checked": False, "none_checked": False}

    def initialize_ui(self):
        # initialize the user interface for the right group box
        right_layout = QVBoxLayout(self)
        right_layout.setContentsMargins(10, 10, 10, 10)
        right_layout.setSpacing(10)

        topLayout = QHBoxLayout()

        self.terminalButton = create_transparent_button(self, "terminal.png", "")
        self.helpButton = create_transparent_button(self, "help.png", "")
        self.settingsButton = create_transparent_button(self, "settings.png", "")

        self.terminalButton.setCursor(QCursor(Qt.PointingHandCursor))
        self.helpButton.setCursor(QCursor(Qt.PointingHandCursor))
        self.settingsButton.setCursor(QCursor(Qt.PointingHandCursor))

        button_size = QSize(64, 64)
        self.terminalButton.setFixedSize(button_size)
        self.helpButton.setFixedSize(button_size)
        self.settingsButton.setFixedSize(button_size)

        self.terminalButton.clicked.connect(self.go_back_to_home)

        self.settingsButton.clicked.connect(self.show_settings_window)

        self.helpButton.clicked.connect(self.show_help_window)

        buttonHolder = QWidget(self)
        buttonLayout = QHBoxLayout(buttonHolder)
        buttonLayout.setContentsMargins(0, 0, 0, 0)
        buttonLayout.setSpacing(10)
        buttonLayout.addWidget(self.terminalButton)
        buttonLayout.addWidget(self.helpButton)
        buttonLayout.addWidget(self.settingsButton)

        self.terminalButton.setStyleSheet("""
        QWidget {
            background-color: #202020;
            border: 2px solid #FF8956;
            border-radius: 10px;
        }
        QWidget:hover {
            background-color: #282828;
        }
        QWidget:focus {
            background-color: #101010;
        }
        """)

        self.helpButton.setStyleSheet("""
        QWidget {
            background-color: #202020;
            border: 2px solid #FF8956;
            border-radius: 10px;
        }
        QWidget:hover {
            background-color: #282828;
        }
        QWidget:focus {
            background-color: #101010;
        }
        """)

        self.settingsButton.setStyleSheet("""
        QWidget {
            background-color: #202020;
            border: 2px solid #FF8956;
            border-radius: 10px;
        }
        QWidget:hover {
            background-color: #282828;
        }
        QWidget:focus {
            background-color: #101010;
        }
        """)

        topLayout.addStretch()
        topLayout.addWidget(buttonHolder)
        topLayout.setAlignment(buttonHolder, Qt.AlignRight)

        right_layout.addLayout(topLayout)

        self.commandInfoBox = QTextEdit(self)
        self.commandInfoBox.setObjectName("textEdit_3")
        self.commandInfoBox.setEnabled(True)
        self.commandInfoBox.setReadOnly(True)
        self.commandInfoBox.setStyleSheet("""
        QTextEdit {
            background-color: transparent;
            border: none;
            color: #F5F3F1;
            font: 15pt "Inter_FXH";
            font-weight: 500;
        }
        QTextEdit:focus {
            outline: none;
        }
        """)
        self.commandInfoBox.setHtml("""
        <div style='text-align: center;'>
            <span style='color: #FF8956;'></span>
        </div>
        """)
        right_layout.addWidget(self.commandInfoBox)


        search_and_table_layout = QVBoxLayout()
        search_and_table_layout.setSpacing(5)

        # search bar
        self.searchBar = QLineEdit(self)
        self.searchBar.setPlaceholderText("Search")
        self.searchBar.setFixedHeight(30)
        self.searchBar.setStyleSheet("""
            QLineEdit {
                background-color: #353535;
                border: 1px solid #FF8956;
                border-radius: 5px;
                color: #FF8956;
                padding: 5px;
                font: 12pt "Inter_FXH";
            }
            QLineEdit:focus {
                outline: none;
            }
            QLineEdit::placeholder {
                color: #FF8956;
            }
        """)
        self.searchBar.textChanged.connect(self.filter_table)
        search_and_table_layout.addWidget(self.searchBar)

        # create and configure the output table
        self.outputTable = QTableWidget(self)
        self.outputTable.setObjectName("outputTable")
        self.outputTable.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        self.outputTable.horizontalHeader().setStretchLastSection(True)
        self.outputTable.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)  # Allow resizing columns
        self.outputTable.setStyleSheet("""
        QTableWidget {
            background-color: #353535;
            border: 1px solid #FF8956;
            border-radius: 10px;
            font: 14pt "Inter_FXH";
            color: white;
            font-weight: 500;
        }
        QHeaderView::section {
            background-color: #454545;
            color: #FF8956;
        }
        """)
        self.outputTable.setSortingEnabled(True)
        self.outputTable.horizontalHeader().sectionClicked.connect(self.handle_header_click)
        self.outputTable.setSelectionBehavior(QTableWidget.SelectRows)
        self.outputTable.setSelectionMode(QTableWidget.SingleSelection)
        self.outputTable.setEditTriggers(QTableWidget.NoEditTriggers)
        self.outputTable.clicked.connect(self.row_clicked)
        search_and_table_layout.addWidget(self.outputTable)

        # progress bar for the output table
        self.outputProgressBar = QProgressBar(self)
        self.outputProgressBar.setRange(0, 100)
        self.outputProgressBar.setValue(0)
        self.outputProgressBar.setVisible(False)
        search_and_table_layout.addWidget(self.outputProgressBar)

        right_layout.addLayout(search_and_table_layout)

        right_layout.addWidget(self.create_spacer(10, ''))

        # create and configure the export and filter buttons
        self.exportButton = QPushButton(self)
        self.exportButton.setFixedSize(330, 50)
        self.exportButton.setCursor(QCursor(Qt.PointingHandCursor))
        setup_button_style(self.exportButton, "Export to...")
        self.exportButton.setSizePolicy(QSizePolicy.Fixed, QSizePolicy.Fixed)
        self.exportButton.setStyleSheet("""
            QPushButton {
                background-color: #FF8956; 
                border: 2px solid black; 
                border-radius: 8px; 
                color: black;
            }

            QPushButton:hover {
                background-color: #FA7B43;
            }

            QPushButton:pressed {
                background-color: #FC6a2B;
            }
        """)

        self.exportButton.clicked.connect(self.check_memdump_path)

        self.filterButton = QPushButton(self)
        setup_button_style(self.filterButton, "Filter")
        self.filterButton.setSizePolicy(QSizePolicy.Fixed, QSizePolicy.Fixed)
        self.filterButton.setFixedSize(330, 50)
        self.filterButton.setCursor(QCursor(Qt.PointingHandCursor))
        self.filterButton.clicked.connect(self.show_filter_window)
        self.filterButton.setStyleSheet("""
            QPushButton {
                background-color: #FF8956; 
                border: 2px solid black; 
                border-radius: 8px; 
                color: black;
            }

            QPushButton:hover {
                background-color: #FA7B43;
            }

            QPushButton:pressed {
                background-color: #FC6a2B;
            }
        """)

        button_layout = QHBoxLayout()
        button_layout.addSpacerItem(QSpacerItem(10, 10, QSizePolicy.Expanding, QSizePolicy.Minimum))
        button_layout.addWidget(self.filterButton)
        button_layout.addWidget(self.exportButton)
        button_layout.addSpacerItem(QSpacerItem(10, 10, QSizePolicy.Expanding, QSizePolicy.Minimum))
        right_layout.addLayout(button_layout)

        right_layout.addWidget(self.create_spacer(10, ''))

        right_layout.setStretchFactor(self.commandInfoBox, 1)
        right_layout.setStretchFactor(search_and_table_layout, 8)
        right_layout.setStretchFactor(button_layout, 1)

        self.setLayout(right_layout)

    def filter_table(self, text):
        # filter the table based on the search text
        for i in range(self.outputTable.rowCount()):
            for j in range(self.outputTable.columnCount()):
                item = self.outputTable.item(i, j)
                if item and text.lower() in item.text().lower():
                    self.outputTable.showRow(i)
                    break
            else:
                self.outputTable.hideRow(i)

    def create_spacer(self, height, color):
        spacer = QWidget()
        spacer.setFixedHeight(height)
        spacer.setStyleSheet(f"background-color: {color};")
        return spacer

    def display_output(self, headers, data):
        # display the output data in the table
        print("Displaying output in table")
        self.headers = headers
        self.data = data

        self.outputTable.setColumnCount(len(headers))
        self.outputTable.setHorizontalHeaderLabels(headers)
        self.outputTable.setRowCount(len(data))

        for row_idx, row_data in enumerate(data):
            for col_idx, col_data in enumerate(row_data):
                item = CustomTableWidgetItem(str(col_data))
                item.setTextAlignment(Qt.AlignCenter)
                self.outputTable.setItem(row_idx, col_idx, item)

        for col in range(len(headers)):
            self.sort_orders[col] = Qt.DescendingOrder

    def handle_header_click(self, logicalIndex):
        # sorting for the table header
        current_order = self.sort_orders.get(logicalIndex, Qt.DescendingOrder)
        self.outputTable.sortItems(logicalIndex, current_order)
        self.sort_orders[logicalIndex] = Qt.AscendingOrder if current_order == Qt.DescendingOrder else Qt.DescendingOrder

    def row_clicked(self, index):
        # emit signal with the data of the selected row and handle PID selection
        selected_row = index.row()
        data = []
        for column in range(self.outputTable.columnCount()):
            data.append(self.outputTable.item(selected_row, column).text())
        self.row_selected_signal.emit(data)

        # handle PID selection
        if 'PID' in self.headers:
            pid_index = self.headers.index('PID')
            selected_pid = self.outputTable.item(selected_row, pid_index).text()
            self.pid_selected_signal.emit(selected_pid)
            print(f"Selected PID: {selected_pid}")  # Print selected PID to the console

    def update_command_info(self, command):
        # update the command info box on top with the executed command
        self.commandInfoBox.setHtml(f"""
        <div style='text-align: center;'>
            <span style='color: #FF8956;'>Executing command:</span>
            <span id='commandText'>{command}</span>
        </div>
        """)

    def clear_output(self):
        # clear button for clearing the output table and command info box
        print("Clearing output")
        self.commandInfoBox.setHtml("""
        <div style='text-align: center;'>
            <span style='color: #FF8956;'></span>
        </div>
        """)
        self.outputTable.setRowCount(0)
        self.outputTable.setColumnCount(0)
        self.outputProgressBar.setVisible(False)
        self.outputProgressBar.setValue(0)

    def open_help(self):
        # open the help URL in the web browser
        webbrowser.open("https://github.com/volatilityfoundation/volatility/wiki/Command-Reference")


    def get_settings_values(self):
        try:
            config = configparser.ConfigParser()
            config.read('settings.ini')

            theme = config['DEFAULT'].get('Theme', 'Light')
            text_size = config['DEFAULT'].get('TextSize', '12')
            text_style = config['DEFAULT'].get('TextStyle', 'Normal')
            memdump_path = config['DEFAULT'].get('MemdumpPath', '')
            file_type = config['DEFAULT'].get('FileType', 'none')
            upload = config['DEFAULT'].get('Upload', 'none')

            return theme, text_size, text_style, memdump_path, file_type, upload
        except Exception as e:
            print(f"Error loading settings: {e}")




    def check_memdump_path(self):
        try:
            theme, text_size, text_style, memdump_path, file_type, upload = self.get_settings_values()
            if memdump_path:
                self.eksport_to_file()
            else:
                self.export_data()
        except Exception as e:
            print(f"Error while checking memdump path: {e}")


    def eksport_to_file(self):
        try:
            print("Starting export process")

            # Get settings values
            theme, text_size, text_style, memdump_path, file_type, upload = self.get_settings_values()
            print(f"Memdump path from settings: {memdump_path}")
            print(f"File type from settings: {file_type}")

            if not memdump_path:
                print("Error: Memdump path is missing in settings.")
                return

            if not os.path.isdir(memdump_path):
                print("Error: Memdump path does not exist or is not a directory.")
                return

            if not os.access(memdump_path, os.W_OK):
                print("Error: No write access to the specified directory.")
                return

            saved_path = memdump_path
            print(f"Saved file path: {saved_path}")

            options = QFileDialog.Options()

            if file_type == "none":
                filePath, _ = QFileDialog.getSaveFileName(
                    self, "Save File", saved_path,
                    "PDF Files (*.pdf);;CSV Files (*.csv);;Excel Files (*.xls);;Text Files (*.txt);;Word Files (*.doc);;JSON Files (*.json)",
                    options=options)
            else:
                options |= QFileDialog.DontConfirmOverwrite
                filePath, _ = QFileDialog.getSaveFileName(
                    self, "Save File", saved_path,
                    f"{file_type.capitalize()} Files (*{file_type});;",
                    options=options)
            print(f"File path selected: {filePath}")

            export_function = None
            if file_type == ".pdf":
                export_function = self.export_to_pdf
            elif file_type == ".csv":
                export_function = self.export_to_csv
            elif file_type == ".xls":
                export_function = self.export_to_excel
            elif file_type == ".txt":
                export_function = self.export_to_txt
            elif file_type == ".doc":
                if DOCX_AVAILABLE:
                    export_function = self.export_to_doc
                else:
                    print("DOCX export not available. Please install python-docx.")
            elif file_type == ".json":
                export_function = self.export_to_json
            else:
                print(f"Error: Unsupported file type: {file_type}")

            if export_function is not None:
                export_function(filePath)
            else:
                print("Error: No export function found for the specified file type.")

        except Exception as e:
            print(f"Unexpected error: {e}")

    # Export the displayed data to a file
    def export_data(self):
        print("Starting export process")
        options = QFileDialog.Options()
        filePath, _ = QFileDialog.getSaveFileName(self, "Save File", "",
                                                  "PDF Files (*.pdf);;CSV Files (*.csv);;Excel Files (*.xls);;Text Files (*.txt);;Word Files (*.doc);;JSON Files (*.json)",
                                                  options=options)
        print(f"File path selected: {filePath}")
        if filePath:
            ext = os.path.splitext(filePath)[1].lower()
            print(f"File extension: {ext}")
            try:
                if ext == ".pdf":
                    self.export_to_pdf(filePath)
                elif ext == ".csv":
                    self.export_to_csv(filePath)
                elif ext == ".xls":
                    self.export_to_excel(filePath)
                elif ext == ".txt":
                    self.export_to_txt(filePath)
                elif ext == ".doc":
                    if DOCX_AVAILABLE:
                        self.export_to_doc(filePath)
                    else:
                        print("DOCX export not available. Please install python-docx.")
                elif ext == ".json":
                    self.export_to_json(filePath)
            except Exception as e:
                print(f"Error during export: {e}")




    def export_to_pdf(self, filePath):
        print(f"Exporting to PDF: {filePath}")
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        line_height = pdf.font_size * 2
        page_width = pdf.w - 2 * pdf.l_margin

        col_width = page_width / len(self.headers)

        for header in self.headers:
            pdf.cell(col_width, line_height, header, border=1)
        pdf.ln(line_height)

        for row in self.data:
            for item in row:
                pdf.cell(col_width, line_height, str(item), border=1)
            pdf.ln(line_height)

        pdf.output(filePath)

    def export_to_csv(self, filePath):
        print(f"Exporting to CSV: {filePath}")
        df = pd.DataFrame(self.data, columns=self.headers)
        df.to_csv(filePath, index=False)

    def export_to_excel(self, filePath):
        print(f"Exporting to Excel: {filePath}")
        df = pd.DataFrame(self.data, columns=self.headers)
        df.to_excel(filePath, index=False)

    def export_to_txt(self, filePath):
        print(f"Exporting to TXT: {filePath}")
        with open(filePath, 'w') as file:
            file.write("\t".join(self.headers) + "\n")
            for row in self.data:
                file.write("\t".join(map(str, row)) + "\n")

    def export_to_doc(self, filePath):
        print(f"Exporting to DOC: {filePath}")
        try:
            from docx import Document
            doc = Document()
            table = doc.add_table(rows=1, cols=len(self.headers))
            hdr_cells = table.rows[0].cells
            for i, header in enumerate(self.headers):
                hdr_cells[i].text = header

            for row in self.data:
                row_cells = table.add_row().cells
                for i, item in enumerate(row):
                    row_cells[i].text = str(item)

            doc.save(filePath)
        except ImportError:
            print("docx module not installed")
        except Exception as e:
            print(f"Error exporting to DOC: {e}")

    def export_to_json(self, filePath):
        print(f"Exporting to JSON: {filePath}")
        df = pd.DataFrame(self.data, columns=self.headers)
        df.to_json(filePath, orient='records', lines=True)

    def show_help_window(self):
        self.settings_window = SettingsWindowGUI()
        self.settings_window.show()

    def show_settings_window(self):
        self.settings_window = SettingsWindowGUI()
        self.settings_window.show()

    def show_help_window(self):
        self.help_window_GUi = helpWindowGui()
        self.help_window_GUi.show()

    def go_back_to_home(self):
        self.back_to_home_signal.emit()

    def apply_filter(self):
        try:
            filter_criteria = self.searchBar.text().strip().lower()
            for i in range(self.outputTable.rowCount()):
                row_match = False
                for j in range(self.outputTable.columnCount()):
                    item = self.outputTable.item(i, j)
                    if item and filter_criteria in item.text().strip().lower():
                        row_match = True
                        break
                self.outputTable.setRowHidden(i, not row_match)
            self.save_filter_settings()
            self.filter_dialog.close()
        except Exception as e:
            print(f"Error applying filter: {e}")

    def toggle_column(self, column_index, checkbox):
        self.outputTable.setColumnHidden(column_index, not checkbox.isChecked())

    def save_filter_settings(self):
        item_vars = {header: self.filter_dialog.findChild(QCheckBox, header) for header in self.headers}
        self.filter_settings["columns"] = {header: checkbox.isChecked() for header, checkbox in item_vars.items()}
        self.filter_settings["all_checked"] = self.filter_dialog.findChild(QCheckBox, "All").isChecked()
        self.filter_settings["none_checked"] = self.filter_dialog.findChild(QCheckBox, "None").isChecked()
        # Save filter settings to a file or database
        with open("filter_settings.json", "w") as f:
            json.dump(self.filter_settings, f)

    def show_filter_window(self):
        self.filter_dialog = QDialog(self)
        self.filter_dialog.setWindowTitle("Filter")
        self.filter_dialog.setStyleSheet("background-color: black; color: white; border: 1px solid #FF8956;")
        self.filter_dialog.setFixedSize(300, 400)

        layout = QVBoxLayout(self.filter_dialog)

        scroll = QScrollArea(self.filter_dialog)
        scroll.setWidgetResizable(True)
        scroll.setStyleSheet("background-color: black;")

        scroll_widget = QWidget()
        scroll_layout = QVBoxLayout(scroll_widget)
        scroll_widget.setLayout(scroll_layout)

        # "All" and "None" checkboxes
        all_checkbox = QCheckBox("All")
        none_checkbox = QCheckBox("None")
        all_checkbox.setObjectName("All")
        none_checkbox.setObjectName("None")
        all_checkbox.setStyleSheet("color: white; border: none;")
        none_checkbox.setStyleSheet("color: white; border: none;")
        all_checkbox.stateChanged.connect(lambda state: self.toggle_all_checkboxes(scroll_widget, state, all_checkbox, none_checkbox))
        none_checkbox.stateChanged.connect(lambda state: self.toggle_all_checkboxes(scroll_widget, not state, none_checkbox, all_checkbox))

        # Load saved state for "All" and "None" checkboxes
        all_checkbox.setChecked(self.filter_settings.get("all_checked", False))
        none_checkbox.setChecked(self.filter_settings.get("none_checked", False))

        scroll_layout.addWidget(all_checkbox)
        scroll_layout.addWidget(none_checkbox)

        item_vars = {}

        # get stored filter settings or default to all True
        filter_settings = self.filter_settings.get("columns", {header: True for header in self.headers})
        for idx, header in enumerate(self.headers):
            checkbox = QCheckBox(header)
            checkbox.setObjectName(header)
            checkbox.setStyleSheet("color: white; border: none;")

            # Set checkbox state based on stored settings
            checkbox.setChecked(filter_settings.get(header, True))
            checkbox.stateChanged.connect(partial(self.toggle_column, idx, checkbox))
            scroll_layout.addWidget(checkbox)
            item_vars[header] = checkbox

        scroll.setWidget(scroll_widget)

        layout.addWidget(scroll)

        button_layout = QHBoxLayout()
        apply_button = QPushButton("Apply")
        apply_button.setStyleSheet("background-color: #FF8956; color: black;")
        apply_button.clicked.connect(self.apply_filter)
        button_layout.addWidget(apply_button)

        layout.addLayout(button_layout)

        self.filter_dialog.setLayout(layout)
        self.filter_dialog.exec_()

    def toggle_all_checkboxes(self, parent_widget, check_state, source_checkbox, other_checkbox):
        other_checkbox.blockSignals(True)
        other_checkbox.setChecked(False)
        other_checkbox.blockSignals(False)
        for checkbox in parent_widget.findChildren(QCheckBox):
            if checkbox not in [source_checkbox, other_checkbox]:
                checkbox.setChecked(check_state)

    def show_progress_bar(self):
        self.outputProgressBar.setVisible(True)

    def update_progress_bar(self, value):
        self.outputProgressBar.setValue(value)
